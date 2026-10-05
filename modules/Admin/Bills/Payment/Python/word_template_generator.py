import docx
import re
import sys
from docx.shared import Inches

def find_and_replace_in_doc(doc, replacements):
    for p in doc.paragraphs:
        for old, new in replacements.items():
            if old in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if old in inline[i].text:
                        text = inline[i].text.replace(old, new)
                        inline[i].text = text

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                find_and_replace_in_doc(cell, replacements)

def main():
    try:
        memo_file = sys.argv[1]
        forwarding_file = sys.argv[2]
        output_file = sys.argv[3]

        replacements = {
            "1st R/A Bill": "${office_bill_no}",
            "M/o DDA Dwarka Golf Course, Sector - 24.": "${name_of_work}",
            "Deployment of Manpower for various purposes at DDA Golf Course Dwarka (DGCD).": "${sub_head}",
            "M/s. Paradise Enterprises": "${agency_name}",
            "4,92,651": "${gross_amount_rounded}",
            "4,48,311": "${net_amount_rounded}",
            "Four Lakh Forty-Eight Thousand Three Hundred Eleven": "${net_amount_words}",
            "AESPB5470P": "${pan_no}",
            "07AESPB5470P1ZV": "${gst_no}",
            "91,20,905": "${tendered_amount}",
            "3,94,662": "${gross_amount_rounded}",
            "3,39,409": "${net_amount_rounded}",
            "1335.00": "${budget_provision}",
            "90541010003984": "${account_no}",
            "CNRB0002455": "${ifsc}",
            "Canara Bank, Vikas Puri, Delhi": "${bank_name}",
            "AAO (DGCD), DDA": "${signatory_aao_dda}",
            "Secretary (DGCD), DDA": "${signatory_secy_dda}",
        }

        # Create a new document to merge into
        merged_doc = docx.Document()
        
        # Process and append Memo
        memo_doc = docx.Document(memo_file)
        find_and_replace_in_doc(memo_doc, replacements)
        for element in memo_doc.element.body:
            merged_doc.element.body.append(element)
        
        # Add a page break
        merged_doc.add_page_break()

        # Process and append Forwarding
        forwarding_doc = docx.Document(forwarding_file)
        find_and_replace_in_doc(forwarding_doc, replacements)
        for element in forwarding_doc.element.body:
            merged_doc.element.body.append(element)

        merged_doc.save(output_file)
        print(f"Successfully created merged template: {output_file}")

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()
