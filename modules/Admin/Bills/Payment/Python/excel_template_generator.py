import openpyxl
import re
import sys

def find_and_replace_in_sheet(sheet, replacements):
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                for old, new in replacements.items():
                    if re.search(re.escape(str(old)), cell.value, re.IGNORECASE):
                        cell.value = cell.value.replace(str(old), new)

def main():
    try:
        source_file = sys.argv[1]
        output_file = sys.argv[2]
        file_type = sys.argv[3] # 'abstract' or 'officestaff'

        if file_type == 'abstract':
            replacements = {
                "1st R/A Bill": "${office_bill_no}",
                "M/o Dwarka Golf Course, Sector-24.": "${name_of_work}",
                "Providing various staff for office maintenance at Dwarka Golf Course.": "${sub_head}",
                "/DWGC/DDA/2023-24": "${agreement_no}",
                "M/s. High Command": "${agency_name}",
                "01 June 2025 to 30 June 2025": "${bill_period}",
                "395,536": "${total_basic}",
                "7,692": "${shortage_amount}",
                "387,844": "${total_a}",
                "32,501": "${profit_amount}",
                "404,961": "${gross_amount}",
                "368,514": "${net_amount}",
            }
        elif file_type == 'officestaff':
            # Add replacements specific to the Office Staff file here
            replacements = {
                "3rd R/A Bill": "${office_bill_no}",
                # ... add other specific placeholders
            }
        else:
            raise ValueError("Unknown file type for Excel processing")

        workbook = openpyxl.load_workbook(source_file)
        sheet = workbook.active
        
        find_and_replace_in_sheet(sheet, replacements)
        
        workbook.save(output_file)
        print(f"Successfully created template: {output_file}")

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main()
